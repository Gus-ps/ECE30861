"""Generate complete XLSX inputs for five-rule partial WCAG integration only.

1.1.1, 1.3.1, 2.4.2 (related 2.4.6), 2.4.4. Not a six-rule clean claim.
Run explicitly to create NEW fixtures; tests never regenerate saved inputs.
"""
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableColumn


def build(path, seeded):
    book = Workbook()
    data = book.active
    data.title = 'Data'
    data.append(['Name', 'Score'])
    data.append(['Ada', 9])
    data.add_table(Table(displayName='Scores', ref='A1:B2', headerRowCount=0 if seeded else 1,
                         tableColumns=[TableColumn(id=1, name='Name'), TableColumn(id=2, name='Score')]))
    objects = book.create_sheet('Sheet2' if seeded else 'Objects')
    objects.sheet_state = 'hidden'
    chart = BarChart()
    chart.add_data(Reference(data, min_col=2, min_row=2, max_row=2))
    objects.add_chart(chart, 'G1')
    objects.merge_cells('A5:C5')
    if seeded:
        objects['A5'] = 'Quarterly results'
    objects['E5'] = 'click here' if seeded else 'Annual accessibility report'
    objects['E5'].hyperlink = 'https://example.invalid/report'
    book.create_sheet('Archive').sheet_state = 'veryHidden'
    book.save(path)  # Fixture generation only. Production never saves.
    book.close()
    if not seeded:
        with ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        root = ET.fromstring(parts['xl/drawings/drawing1.xml'])
        props = root.find('.//{*}cNvPr')
        props.set('title', 'Scores')
        props.set('descr', 'Score for Ada')
        parts['xl/drawings/drawing1.xml'] = ET.tostring(root)
        with ZipFile(path, 'w') as archive:
            for name, content in parts.items():
                archive.writestr(name, content)


if __name__ == '__main__':
    for seeded in (False, True):
        build(Path(__file__).parent / ('five_rule_seeded.xlsx' if seeded else 'five_rule_clean.xlsx'), seeded)
