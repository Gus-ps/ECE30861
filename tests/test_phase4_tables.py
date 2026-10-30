"""1.3.1 Phase 4 formal tables; REQ-03/04/05/06 and AC-03/04/05/06."""
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableColumn
from accessibility_checker.checks.table_headers import check_table_headers

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


class Phase4FormalTableTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'tables.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = 'Data'
        sheet.append(['Name', 'Amount'])
        sheet.append(['Item', 0])
        table = Table(id=1, displayName='Budget', ref='A1:B2', headerRowCount=0,
                      tableColumns=[TableColumn(id=1, name='Name'), TableColumn(id=2, name='Amount')])
        sheet.add_table(table)
        book.save(self.path)
        book.close()

    def edit(self, part, action):
        with ZipFile(self.path) as archive:
            parts = {p: archive.read(p) for p in archive.namelist()}
        root = ET.fromstring(parts[part])
        action(root)
        parts[part] = ET.tostring(root)
        with ZipFile(self.path, 'w') as archive:
            for p, data in parts.items():
                archive.writestr(p, data)

    def test_safe_disabled_exact_identity_stability_and_bytes(self):
        before = self.path.read_bytes()
        result = check_table_headers(self.path)
        self.assertEqual(len(result), 1)
        item = result[0]
        self.assertEqual((item['ruleId'], item['category'], item['classification'], item['target']),
                         ('TABLE_HEADER', 'TABLE_STRUCTURE', 'REQUIRES_USER_INPUT', 'xl/tables/table1.xml#tableId=1'))
        payload = [hashlib.sha256(before).hexdigest(), '2.0.0', 'TABLE_HEADER', 'Data', item['target']]
        self.assertEqual(item['findingId'], hashlib.sha256(json.dumps(payload, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest())
        self.assertEqual(result, check_table_headers(self.path))
        self.assertEqual(before, self.path.read_bytes())

    def test_enabled_default_and_invalid_count(self):
        self.edit('xl/tables/table1.xml', lambda r: r.attrib.pop('headerRowCount'))
        self.assertEqual(check_table_headers(self.path), [])
        self.edit('xl/tables/table1.xml', lambda r: r.set('headerRowCount', '2'))
        self.assertEqual(check_table_headers(self.path)[0]['classification'], 'MANUAL_REVIEW')

    def test_missing_id_and_metadata_mismatch(self):
        self.edit('xl/tables/table1.xml', lambda r: r.attrib.pop('id'))
        result = check_table_headers(self.path)
        self.assertEqual(result[0]['target'], 'xl/tables/table1.xml#xml=/table[1]')
        self.assertEqual(result[0]['classification'], 'MANUAL_REVIEW')
        self.assertTrue(result.limitations)

    def test_totals_query_extension_and_header_shape_are_manual(self):
        for key, value in [('totalsRowCount', '1'), ('connectionId', '4'), ('ref', 'A1:B1')]:
            with self.subTest(key=key):
                self.setUp()
                self.edit('xl/tables/table1.xml', lambda r: r.set(key, value))
                self.assertEqual(check_table_headers(self.path)[0]['classification'], 'MANUAL_REVIEW')

    def test_formula_rich_blank_duplicate_headers(self):
        for mode in ['formula', 'rich', 'blank', 'duplicate']:
            with self.subTest(mode=mode):
                self.setUp()
                def mutate(root):
                    cell = root.find(NS + 'sheetData/' + NS + 'row/' + NS + 'c')
                    cell.clear()
                    cell.set('r', 'A1')
                    if mode == 'formula':
                        ET.SubElement(cell, NS+'f').text = '1'
                    else:
                        cell.set('t', 'inlineStr')
                        body = ET.SubElement(cell, NS+'is')
                        if mode == 'rich':
                            body = ET.SubElement(body, NS+'r')
                        ET.SubElement(body, NS+'t').text = {'blank': ' \u00a0', 'duplicate': ' amount ', 'rich': 'Name'}[mode]
                self.edit('xl/worksheets/sheet1.xml', mutate)
                result = check_table_headers(self.path)
                self.assertEqual(result[0]['classification'], 'MANUAL_REVIEW')

    def test_hidden_and_merge(self):
        self.edit('xl/workbook.xml', lambda r: r.find(NS+'sheets/'+NS+'sheet').set('state', 'veryHidden'))
        self.assertEqual(check_table_headers(self.path)[0]['classification'], 'REQUIRES_USER_INPUT')
        def merge(root):
            ET.SubElement(ET.SubElement(root, NS+'mergeCells'), NS+'mergeCell', ref='A1:B1')
        self.edit('xl/worksheets/sheet1.xml', merge)
        self.assertEqual(check_table_headers(self.path)[0]['classification'], 'MANUAL_REVIEW')

    def test_shared_table_and_duplicate_ids_all_members(self):
        # Clone the physical table and reference it from another worksheet.
        from openpyxl import load_workbook
        book = load_workbook(self.path)
        book.create_sheet('Other').append(['Name', 'Amount'])
        book.save(self.path)
        book.close()
        doc = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
        pkg = 'http://schemas.openxmlformats.org/package/2006/relationships'
        with ZipFile(self.path) as archive:
            parts = {p: archive.read(p) for p in archive.namelist()}
        sheet = ET.fromstring(parts['xl/worksheets/sheet2.xml'])
        ET.SubElement(ET.SubElement(sheet, NS+'tableParts'), NS+'tablePart', {f'{{{doc}}}id': 't'})
        parts['xl/worksheets/sheet2.xml'] = ET.tostring(sheet)
        for target, duplicate in [('../tables/table1.xml', False), ('../tables/table9.xml', True)]:
            rels = ET.Element(f'{{{pkg}}}Relationships')
            ET.SubElement(rels, f'{{{pkg}}}Relationship', Id='t', Type=doc+'/table', Target=target)
            parts['xl/worksheets/_rels/sheet2.xml.rels'] = ET.tostring(rels)
            if duplicate:
                parts['xl/tables/table9.xml'] = parts['xl/tables/table1.xml']
            with ZipFile(self.path, 'w') as archive:
                for name, data in parts.items():
                    archive.writestr(name, data)
            result = check_table_headers(self.path)
            self.assertEqual(len(result), 2)
            self.assertTrue(all(f['classification'] == 'MANUAL_REVIEW' for f in result))
            if duplicate:
                self.assertEqual({f['target'] for f in result}, {'xl/tables/table1.xml#xml=/table[1]', 'xl/tables/table9.xml#xml=/table[1]'})
            else:
                self.assertTrue(any('multiple owning' in reason for reason in result.limitations))

class UnstructuredTableTests(unittest.TestCase):
    def run_cells(self, rows, *, merge=None, auto_filter=None, formal=False):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / 'grid.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = 'Data'
        for reference, value in rows.items():
            sheet[reference] = value
        if merge:
            sheet.merge_cells(merge)
        if auto_filter:
            sheet.auto_filter.ref = auto_filter
        if formal:
            sheet.add_table(Table(displayName='Already', ref='A1:B2'))
        book.save(path)
        book.close()
        return path, check_table_headers(path)

    def test_dense_with_formula_body_and_false_zero(self):
        path, result = self.run_cells({'A1':'Name', 'B1':'Amount', 'A2':False, 'B2':'=1'})
        self.assertEqual([(f['target'], f['classification']) for f in result],
                         [('xl/worksheets/sheet1.xml#range=A1:B2', 'REQUIRES_USER_INPUT')])
        self.assertEqual(result.limitations, [])
        before = path.read_bytes()
        self.assertEqual(result, check_table_headers(path))
        self.assertEqual(before, path.read_bytes())

    def test_sparse_l_shape_and_diagonal_are_not_dense(self):
        _, result = self.run_cells({'A1':'Name', 'B1':'Amount', 'A2':0})
        self.assertEqual(result[0]['classification'], 'MANUAL_REVIEW')
        self.assertTrue(any('not a dense' in x for x in result.limitations))
        _, result = self.run_cells({'A1':'Name', 'B2':0, 'XFD1048576':False})
        self.assertEqual(result, [])
        self.assertEqual(len(result.limitations), 1)

    def test_empty_row_separates_components_and_formal_is_excluded(self):
        _, result = self.run_cells({'A1':'Name', 'B1':'Amount', 'A2':0, 'B2':False,
                                   'A4':'Other', 'B4':'Value', 'A5':0, 'B5':1})
        self.assertEqual([f['target'].split('#')[1] for f in result], ['range=A1:B2','range=A4:B5'])
        _, result = self.run_cells({'A1':'Name', 'B1':'Amount', 'A2':0, 'B2':False}, formal=True)
        self.assertEqual(result, [])

    def test_filter_merge_and_header_formula_are_manual(self):
        cells = {'A1':'Name', 'B1':'Amount', 'A2':0, 'B2':False}
        for option in ({'auto_filter':'A1:B2'}, {'merge':'A1:B1'}):
            _, result = self.run_cells(cells, **option)
            self.assertEqual(result[0]['classification'], 'MANUAL_REVIEW')
            self.assertTrue(result.limitations)
        _, result = self.run_cells(dict(cells, A1='=1'))
        self.assertEqual(result[0]['classification'], 'MANUAL_REVIEW')
