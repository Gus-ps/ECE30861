"""Checks against the fake target: the fixture-free intake checks, and the rejected-fixture
checks run with Team 12's shared pack."""

import base64
import io
import json
import re
import shutil
import zipfile

import httpx
import openpyxl
import pytest

from tester.config import FIXTURES
from tester.suite.checks import CHECKS
from tester.suite.runner import CaseResult
from tester.suite.transfer import upload_files
from tester.suite.workbooks import PADDING_PART, column, structure_password_hash, workbook

from .conftest import make_settings
from .fake_project import FakeProject, multipart
from .test_fixtures_and_record import run as run_suite
from .test_fixtures_and_record import setup as setup_session

FIXTURE_FREE = [
    "P3-HEALTH", "P3-AUTH-UPLOAD", "P3-AUTH-UPLOAD-INVALID", "P3-INPUT-MISSING-FILE", "P3-INPUT-CSV",
    "P3-INPUT-MISSING-EXTENSION", "P3-INPUT-XLS", "P3-INPUT-XLSM", "P3-INPUT-PROTECTED", "P3-INPUT-BYTES-OVER",
    "P3-INPUT-BYTES-AT", "P3-INPUT-SHEETS-OVER", "P3-INPUT-SHEETS-AT", "P3-INPUT-CELLS-OVER", "P3-INPUT-CELLS-AT",
    "P3-NOT-FOUND", "P3-SESSION-DUPLICATE", "P3-CODE-REPORT",
]
SHARED_PACK = FIXTURES / "shared" / "expectations.json"


def results(client):
    body = run_suite(client, setup_session(client).json()["session_id"]).json()
    return {r["test_id"]: r for r in body["results"]}, body


def only(*names):
    return {name: CHECKS[name] for name in names}


def uploads(fake):
    ##  (headers, {field: (filename, bytes)}) for every upload the tester sent.
    return [(headers, multipart(headers["content-type"], body)) for method, path, headers, body in fake.requests
            if method == "POST" and path == "/system/api/v1/workbooks"]


##  ---- Against a conforming target ----

def test_fixture_free_checks_pass_against_a_conforming_target(make_client):
    by_id, body = results(make_client())
    assert [t for t, r in by_id.items() if r["outcome"] == "pass"] == FIXTURE_FREE
    assert {r["outcome"] for t, r in by_id.items() if t not in FIXTURE_FREE} == {"skip"}
    assert body["coverage"]["component"]["covered"] == 4          # C1, C2, C3, C6
    assert body["coverage"]["system"]["covered"] == 3             # AC-01, AC-02, AC-15
    assert by_id["P3-INPUT-XLSM"]["message"].endswith("code MACRO_ENABLED_WORKBOOK_UNSUPPORTED.")


def test_requests_carry_the_intended_inputs(make_client, fake):
    results(make_client())
    health = next(r for r in fake.requests if r[1] == "/system/api/v1/health")
    assert "authorization" not in health[2] and health[2]["x-test-session"]
    sent = uploads(fake)
    names = [next(iter(fields.values()))[0] for _, fields in sent]
    assert names == ["intake.xlsx", "intake.xlsx", "intake.xlsx", "totals.csv", "intake", "legacy.xls", "macros.xlsm",
                     "protected.xlsx", "over-limit.xlsx", "at-limit.xlsx", "worksheets-over.xlsx",
                     "worksheets-at.xlsx", "cells-over.xlsx", "cells-at.xlsx"]
    assert all(len(fields) == 1 for _, fields in sent)
    assert "authorization" not in sent[0][0]
    assert re.fullmatch(r"Bearer tester-invalid-[0-9a-f]{32}", sent[1][0]["authorization"])
    assert list(sent[2][1]) == ["attachment"]
    assert all(h["authorization"] == "Bearer tok-a" for h, _ in sent[2:])
    assert [len(f["file"][1]) for _, f in sent[8:10]] == [25_000_001, 25_000_000]


##  ---- Against a target that gets things wrong ----

def test_a_target_that_accepts_everything_fails_every_rejection(make_client, fake):
    accepted = {"workbookId": "wb-1", "versionId": "ver-1", "auditId": "audit-1", "status": "PROCESSING"}
    fake.project_reply = lambda request: (202, accepted) if request.method == "POST" else None
    by_id, _ = results(make_client())
    for test_id in ["P3-AUTH-UPLOAD", "P3-INPUT-CSV", "P3-INPUT-XLSM", "P3-INPUT-BYTES-OVER", "P3-INPUT-CELLS-OVER"]:
        assert by_id[test_id]["outcome"] == "fail" and "accepting the upload" in by_id[test_id]["message"]
    ##  Its at-limit uploads are accepted, but their audits cannot be read, so the wait fails them.
    at_limit = by_id["P3-INPUT-BYTES-AT"]
    assert at_limit["outcome"] == "fail" and "was accepted (202), but: Expected 200" in at_limit["message"]


##  ---- Boundary uploads wait for their audits (Team 9 feedback, 2026-11-01) ----

BOUNDARY = ["at-limit.xlsx", "worksheets-at.xlsx", "cells-at.xlsx"]


def test_each_boundary_audit_finishes_before_the_next_case_starts(make_client, fake, tmp_path):
    fake.project.processing_polls = 2
    by_id, _ = results(make_client(settings=make_settings(tmp_path, FAST_POLL)))
    for test_id in ("P3-INPUT-BYTES-AT", "P3-INPUT-SHEETS-AT", "P3-INPUT-CELLS-AT"):
        assert by_id[test_id]["outcome"] == "pass" and "its audit finished as" in by_id[test_id]["message"]
    ##  After each boundary upload, its audit is read until finished (2 PROCESSING reads, then the
    ##  result) and nothing else is sent to the project API in between.
    sequence = [r for r in fake.requests if "/api/v1/" in r[1]]
    for name in BOUNDARY:
        start = next(i for i, (method, _p, headers, body) in enumerate(sequence) if method == "POST"
                     and next(iter(multipart(headers["content-type"], body).values()))[0] == name)
        following = sequence[start + 1:start + 4]
        assert [m for m, *_ in following] == ["GET"] * 3 and all("/audits/" in p for _m, p, *_ in following), name
        assert len({p for _m, p, *_ in following}) == 1, name
    assert len(audit_reads(fake)) == 3 * len(BOUNDARY)


def test_a_boundary_audit_still_processing_at_the_deadline_fails_the_case(make_client, fake, tmp_path):
    fake.project.processing_polls = 10 ** 6
    settings = make_settings(tmp_path, {"timing": {"poll_interval_seconds": 0.01, "audit_poll_seconds": 0.1}})
    result = results(make_client(checks=only("upload_worksheets_at"), settings=settings))[0]["P3-INPUT-SHEETS-AT"]
    assert result["outcome"] == "fail"
    assert result["message"] == ("A workbook with 50 worksheets was accepted (202), but: The audit of a workbook with "
                                 "50 worksheets was still PROCESSING after 0.1 s.")
    assert len(uploads(fake)) == 1          # never re-uploaded


def test_a_failed_boundary_audit_fails_the_case_without_a_retry(make_client, fake, tmp_path):
    mutate_audits(fake, lambda b: b.update(status="PROCESSING_FAILED", findings=[], error={
        "code": "AUDIT_FAILED", "message": "No.", "requestId": "r1", "details": {}}))
    settings = make_settings(tmp_path, FAST_POLL)
    result = results(make_client(checks=only("upload_worksheets_at"), settings=settings))[0]["P3-INPUT-SHEETS-AT"]
    assert result == {"test_id": "P3-INPUT-SHEETS-AT", "outcome": "fail", "message":
                      "A workbook with 50 worksheets was accepted (202), but its audit ended PROCESSING_FAILED "
                      "(AUDIT_FAILED)."}
    assert len(uploads(fake)) == 1 and len(audit_reads(fake)) == 1


@pytest.mark.parametrize("reply, problem", [
    ({"code": "X", "message": "No.", "requestId": "r1"}, "details: Field required"),
    ({"code": "X", "message": "No.", "requestId": "r1", "details": {}, "trace": "..."}, "trace: Extra inputs"),
    ({"code": "X", "message": "No.", "requestId": "has spaces", "details": {}}, "requestId"),
    ({"code": " ", "message": "No.", "requestId": "r1", "details": {}}, "code"),
    (b"<html>401</html>", "body: Input should be a valid dictionary"),
])
def test_rejection_body_must_be_a_strict_project_error(make_client, fake, reply, problem):
    fake.project_reply = lambda request: (401, reply)
    by_id, _ = results(make_client(checks=only("upload_without_token")))
    result = by_id["P3-AUTH-UPLOAD"]
    assert result["outcome"] == "fail" and "does not match ProjectError" in result["message"]
    assert problem in result["message"], result["message"]


def test_xlsm_code_is_compared_exactly(make_client, fake):
    fake.project_reply = lambda request: (422, {"code": "INVALID_WORKBOOK", "message": "No.", "requestId": "r1",
                                                "details": {}})
    result = results(make_client(checks=only("upload_xlsm", "upload_protected")))[0]
    assert result["P3-INPUT-XLSM"]["outcome"] == "fail"
    assert "code INVALID_WORKBOOK, expected MACRO_ENABLED_WORKBOOK_UNSUPPORTED" in result["P3-INPUT-XLSM"]["message"]
    assert result["P3-INPUT-PROTECTED"]["outcome"] == "pass"           # other codes are not compared


@pytest.mark.parametrize("reply, outcome, text", [
    ((200, {"status": "AVAILABLE", "apiVersion": "1.0", "rulesetVersion": "2.0.0"}), "fail", "ruleset 2.0.0, expected 1.0.0"),
    ((200, {"status": "AVAILABLE", "apiVersion": 1.0, "rulesetVersion": "1.0.0"}), "fail", "apiVersion"),
    ((200, {"status": "AVAILABLE", "apiVersion": "1.0", "rulesetVersion": "1.0"}), "fail", "rulesetVersion"),
    ((401, {"code": "X", "message": "No.", "requestId": "r1", "details": {}}), "fail", "received 401"),
    ((503, {"code": "X", "message": "Down.", "requestId": "r1", "details": {}}), "error", "unexpected 503"),
])
def test_health(make_client, fake, reply, outcome, text):
    fake.project_reply = lambda request: reply
    result = results(make_client(checks=only("health")))[0]["P3-HEALTH"]
    assert result["outcome"] == outcome and text in result["message"], result


@pytest.mark.parametrize("reply, text", [
    ((413, {"code": "X", "message": "Too big.", "requestId": "r1", "details": {}}), "Expected 202"),
    ((202, {"workbookId": "same", "versionId": "same", "auditId": "a1", "status": "PROCESSING"}), "not distinct"),
    ((202, {"workbookId": "w1", "versionId": "v1", "auditId": "a1", "status": "AUDIT_COMPLETE"}), "status"),
])
def test_at_limit_needs_a_valid_upload_response(make_client, fake, reply, text):
    fake.project_reply = lambda request: reply
    result = results(make_client(checks=only("upload_at_limit")))[0]["P3-INPUT-BYTES-AT"]
    assert result["outcome"] == "fail" and text in result["message"], result


def test_rejection_5xx_is_a_system_error_with_no_coverage(make_client, fake):
    fake.project_reply = lambda request: (500, {"code": "X", "message": "Boom.", "requestId": "r1", "details": {}})
    by_id, body = results(make_client(checks=only("upload_csv")))
    assert by_id["P3-INPUT-CSV"]["outcome"] == "error" and by_id["P3-INPUT-CSV"]["source"] == "system"
    assert body["coverage"]["component"]["covered"] == 0


def test_unknown_resources_without_uploads_runs_only_version_unknown(make_client):
    by_id, _ = results(make_client(checks=only("unknown_resources")))
    assert by_id["P3-NOT-FOUND"]["outcome"] == "pass"
    assert by_id["P3-VERSION-MISMATCH"]["outcome"] == "skip"
    assert "prerequisite P3-UPLOAD-CLEAN" in by_id["P3-VERSION-MISMATCH"]["message"]


##  ---- Rejected fixtures from the shared pack ----

REJECTED_CASES = {"P3-INPUT-CORRUPT": "corrupt", "P3-INPUT-ZIP-TRAVERSAL": "zip_traversal",
                  "P3-INPUT-ZIP-DUPLICATE": "zip_duplicate", "P3-INPUT-XML-DTD": "xml_dtd",
                  "P3-INPUT-MACRO-PARTS": "macro_parts", "P3-INPUT-ENCRYPTED-ZIP": "encrypted_zip",
                  "P3-INPUT-PROTECTED-SHEET": "protected_sheet", "P3-INPUT-PROTECTED-STRUCTURE": "protected_structure"}


def with_pack(make_client, tmp_path, pack=SHARED_PACK, checks=None):
    return make_client(checks=checks or only("upload_corrupt", "rejected_fixture"),
                       settings=make_settings(tmp_path, expectations_path=pack))


def test_rejected_fixtures_pass_against_a_conforming_target(make_client, fake, tmp_path):
    by_id, body = results(with_pack(make_client, tmp_path))
    for test_id, key in REJECTED_CASES.items():
        assert by_id[test_id]["outcome"] == "pass", by_id[test_id]
        assert by_id[test_id]["message"] == f"The {key} fixture ({key}.xlsx) returned 422 with a valid project error."
    sent = {fields["file"][0]: fields["file"][1] for _, fields in uploads(fake)}
    assert sent == {f"{key}.xlsx": (SHARED_PACK.parent / f"{key}.xlsx").read_bytes() for key in REJECTED_CASES.values()}
    assert body["coverage"]["component"]["covered"] == 1 and body["coverage"]["system"]["covered"] == 1  # C2, AC-02


def test_rejected_fixtures_fail_when_the_target_accepts_them(make_client, fake, tmp_path):
    accepted = {"workbookId": "wb-1", "versionId": "ver-1", "auditId": "audit-1", "status": "PROCESSING"}
    fake.project_reply = lambda request: (202, accepted)
    by_id, _ = results(with_pack(make_client, tmp_path))
    for test_id in REJECTED_CASES:
        assert by_id[test_id]["outcome"] == "fail" and "accepting the upload" in by_id[test_id]["message"]


def test_rejected_fixture_codes_are_compared_only_when_the_pack_names_one(make_client, fake, tmp_path):
    pack = json.loads(SHARED_PACK.read_text(encoding="utf-8"))
    pack["fixtures"]["xml_dtd"]["error_code"] = "XML_DTD_NOT_ALLOWED"
    copy = tmp_path / "pack"
    shutil.copytree(SHARED_PACK.parent, copy)
    (copy / "expectations.json").write_text(json.dumps(pack), encoding="utf-8")
    by_id, _ = results(with_pack(make_client, tmp_path, pack=copy / "expectations.json"))
    assert by_id["P3-INPUT-XML-DTD"]["outcome"] == "fail"
    assert "code INVALID_WORKBOOK, expected XML_DTD_NOT_ALLOWED" in by_id["P3-INPUT-XML-DTD"]["message"]
    assert by_id["P3-INPUT-ZIP-TRAVERSAL"]["outcome"] == "pass"


def test_rejected_fixtures_skip_without_a_pack(make_client, fake):
    by_id, _ = results(make_client(checks=only("upload_corrupt", "rejected_fixture")))
    for test_id, key in REJECTED_CASES.items():
        assert by_id[test_id]["outcome"] == "skip"
        assert by_id[test_id]["message"] == f"Not executed: no fixture pack with a {key} entry was frozen at setup."
    assert uploads(fake) == []


def test_an_accepted_entry_for_a_rejection_case_is_a_tester_error(make_client, fake, tmp_path):
    pack = json.loads(SHARED_PACK.read_text(encoding="utf-8"))
    pack["fixtures"]["macro_parts"].update(kind="accepted", http_status=None)
    copy = tmp_path / "pack"
    shutil.copytree(SHARED_PACK.parent, copy)
    (copy / "expectations.json").write_text(json.dumps(pack), encoding="utf-8")
    result = results(with_pack(make_client, tmp_path, pack=copy / "expectations.json"))[0]["P3-INPUT-MACRO-PARTS"]
    assert result["outcome"] == "error" and result["source"] == "tester"
    assert not any(fields["file"][0] == "macro_parts.xlsx" for _, fields in uploads(fake))


@pytest.mark.parametrize("key, code", [
    ("zip_traversal", "INVALID_WORKBOOK"), ("zip_duplicate", "INVALID_WORKBOOK"), ("xml_dtd", "INVALID_WORKBOOK"),
    ("macro_parts", "INVALID_WORKBOOK"), ("encrypted_zip", "INVALID_WORKBOOK"), ("protected_sheet", "PROTECTED_WORKBOOK"),
    ("protected_structure", "PROTECTED_WORKBOOK"), ("corrupt", "INVALID_WORKBOOK"),
    ("clean", None), ("seeded_issues", None), ("unicode", None), ("repeated_ids", None), ("multiple_drawings", None),
])
def test_fake_project_rejects_each_fixture_for_its_own_defect(key, code):
    ##  Guards the fake, not the target: a rejected fixture must not be refused for a reason
    ##  the valid fixtures share, and every valid fixture must still be accepted.
    content = (SHARED_PACK.parent / f"{key}.xlsx").read_bytes()
    status, reply = FakeProject(set()).upload(*multipart_body(f"{key}.xlsx", content))
    assert (status, reply.get("code")) == ((202, None) if code is None else (422, code))


def multipart_body(name, content):
    request = httpx.Request("POST", "https://x/", files=upload_files(name, content))
    return request.headers["content-type"], request.read()


##  ---- Accepted fixtures: upload and audit ----

ACCEPTED_CASES = {"P3-UPLOAD-CLEAN": "upload", "P3-UPLOAD-SEEDED": "upload", "P3-AUDIT-CLEAN": "audit",
                  "P3-AUDIT-SEEDED": "audit", "P3-UPLOAD-UNICODE": "upload", "P3-AUDIT-UNICODE": "audit",
                  "P3-UPLOAD-REPEATED-IDS": "upload", "P3-AUDIT-REPEATED-IDS": "audit",
                  "P3-UPLOAD-MULTIPLE-DRAWINGS": "upload", "P3-AUDIT-MULTIPLE-DRAWINGS": "audit"}
AUDIT_CHECKS = ("upload_clean", "upload_seeded", "audit_clean", "audit_seeded", "edge_upload", "edge_audit")
FAST_POLL = {"timing": {"poll_interval_seconds": 0.01}}


def audit_client(make_client, tmp_path, profile_changes=None, checks=AUDIT_CHECKS):
    return make_client(checks=only(*checks), settings=make_settings(tmp_path, profile_changes or FAST_POLL,
                                                                    expectations_path=SHARED_PACK))


def audit_reads(fake):
    return [r for r in fake.requests if r[0] == "GET" and "/audits/" in r[1]]


def mutate_audits(fake, change):
    ##  Answer audit reads from the fake project, then apply change(body) to finished audits.
    def reply(request):
        if "/audits/" not in request.url.path:
            return None
        status, body = fake.project.handle(request.method, request.url.path.removeprefix("/system"),
                                           dict(request.headers), request.content)
        if status == 200 and body["status"] != "PROCESSING":
            body = json.loads(json.dumps(body))
            change(body)
        return status, body
    fake.project_reply = reply


def test_accepted_fixtures_pass_against_a_conforming_target(make_client, fake, tmp_path):
    by_id, body = results(audit_client(make_client, tmp_path))
    for test_id in ACCEPTED_CASES:
        assert by_id[test_id]["outcome"] == "pass", by_id[test_id]
    assert by_id["P3-AUDIT-CLEAN"]["message"] == \
        "The clean audit reached AUDIT_COMPLETE with exactly the 0 expected findings in contract order."
    assert by_id["P3-AUDIT-SEEDED"]["message"] == \
        "The seeded_issues audit reached REVIEW_REQUIRED with exactly the 8 expected findings in contract order."
    assert body["coverage"]["component"]["covered"] == 4      # C2, C3, C4, C5
    sent = [fields["file"][0] for _, fields in uploads(fake)]
    assert sent == ["clean.xlsx", "seeded_issues.xlsx", "unicode.xlsx", "repeated_ids.xlsx", "multiple_drawings.xlsx"]


def test_audit_is_polled_until_it_leaves_processing(make_client, fake, tmp_path):
    fake.project.processing_polls = 2
    by_id, _ = results(audit_client(make_client, tmp_path, checks=("upload_clean", "audit_clean")))
    assert by_id["P3-AUDIT-CLEAN"]["outcome"] == "pass"
    assert len(audit_reads(fake)) == 3


def test_audit_still_processing_at_the_poll_budget_fails_without_coverage(make_client, fake, tmp_path):
    fake.project.processing_polls = 10 ** 6
    profile = {"timing": {"poll_interval_seconds": 0.01, "audit_poll_seconds": 0.1}}
    by_id, body = results(audit_client(make_client, tmp_path, profile, checks=("upload_clean", "audit_clean")))
    assert by_id["P3-AUDIT-CLEAN"]["outcome"] == "fail"
    assert "still PROCESSING after 0.1 s" in by_id["P3-AUDIT-CLEAN"]["message"]
    assert body["coverage"]["component"]["covered"] == 2      # the upload's C2, C3; no C4/C5 from the audit


@pytest.mark.parametrize("change, text", [
    (lambda b: b["findings"].pop(), "missing WORKSHEET_NAME REQUIRES_USER_INPUT on 'Sheet1' at xl/worksheets/sheet2.xml#sheetId=2"),
    (lambda b: b["findings"].append({**b["findings"][0], "target": "xl/drawings/drawing9.xml#id=9",
                                     "findingId": "0" * 64}), "unexpected ALT_TEXT"),
    (lambda b: b["findings"].reverse(), "not in contract order"),
    (lambda b: b["findings"][0].update(findingId="0" * 64), "findingId at xl/drawings/drawing3.xml#id=2 is 000"),
    (lambda b: b["findings"][0].update(classification="MANUAL_REVIEW"), "missing ALT_TEXT REQUIRES_USER_INPUT"),
    (lambda b: b.update(auditId="audit-other"), "auditId audit-other is not the upload's"),
    (lambda b: b.update(rulesetVersion="1.0.1"), "rulesetVersion 1.0.1, expected 1.0.0"),
    (lambda b: b.update(status="AUDIT_COMPLETE"), "status AUDIT_COMPLETE, expected REVIEW_REQUIRED"),
    (lambda b: b.update(completedAt=None), "completedAt is null"),
    (lambda b: b.update(completedAt="2000-01-01T00:00:00Z"), "completedAt is earlier than startedAt"),
    (lambda b: b.update(limitations=["Something was skipped."]), "expected no limitations; received 1"),
    (lambda b: b.update(status="PROCESSING_FAILED", findings=[],
                        error={"code": "AUDIT_FAILED", "message": "No.", "requestId": "r1", "details": {}}),
     "processing failed (AUDIT_FAILED), expected REVIEW_REQUIRED"),
    (lambda b: b.update(trace="x"), "does not match AuditResponse (trace: Extra inputs"),
    (lambda b: b["findings"][0].update(message=""), "does not match AuditResponse (findings.0.message"),
    (lambda b: b["findings"][0].update(category="ALT_TEXT"), "does not match AuditResponse (findings.0.category"),
    (lambda b: b.update(startedAt="2026-11-01 10:00:00"), "does not match AuditResponse (startedAt"),
])
def test_seeded_audit_reports_each_kind_of_mismatch(make_client, fake, tmp_path, change, text):
    mutate_audits(fake, change)
    result = results(audit_client(make_client, tmp_path, checks=("upload_seeded", "audit_seeded")))[0]["P3-AUDIT-SEEDED"]
    assert result["outcome"] == "fail" and text in result["message"], result["message"]


def test_repeated_ids_audit_needs_a_disclosure(make_client, fake, tmp_path):
    mutate_audits(fake, lambda b: b.update(limitations=[]))
    by_id, _ = results(audit_client(make_client, tmp_path, checks=("edge_upload", "edge_audit")))
    assert by_id["P3-AUDIT-REPEATED-IDS"]["outcome"] == "fail"
    assert "limitations is empty" in by_id["P3-AUDIT-REPEATED-IDS"]["message"]
    assert by_id["P3-AUDIT-UNICODE"]["outcome"] == "pass"


def test_wrong_findings_still_leave_the_audit_usable_but_a_failed_audit_does_not(make_client, fake, tmp_path):
    def later(case, ctx):
        return CaseResult.passed("ran")

    checks = {**only("upload_seeded", "audit_seeded"), "hash_after_audit": later}
    settings = make_settings(tmp_path, FAST_POLL, expectations_path=SHARED_PACK)
    mutate_audits(fake, lambda b: b["findings"].pop())
    assert results(make_client(checks=checks, settings=settings))[0]["P3-SOURCE-HASH"]["outcome"] == "pass"
    mutate_audits(fake, lambda b: b.update(status="PROCESSING_FAILED", findings=[], error={
        "code": "AUDIT_FAILED", "message": "No.", "requestId": "r1", "details": {}}))
    assert results(make_client(checks=checks, settings=settings))[0]["P3-SOURCE-HASH"]["outcome"] == "skip"


def test_rejected_upload_skips_its_audit(make_client, fake, tmp_path):
    fake.project_reply = lambda request: ((422, {"code": "INVALID_WORKBOOK", "message": "No.", "requestId": "r1",
                                                 "details": {}}) if request.method == "POST" else None)
    by_id, _ = results(audit_client(make_client, tmp_path, checks=("upload_clean", "audit_clean")))
    assert by_id["P3-UPLOAD-CLEAN"]["outcome"] == "fail" and "Expected 202" in by_id["P3-UPLOAD-CLEAN"]["message"]
    assert by_id["P3-AUDIT-CLEAN"]["outcome"] == "skip" and "prerequisite P3-UPLOAD-CLEAN" in by_id["P3-AUDIT-CLEAN"]["message"]
    assert audit_reads(fake) == []


def test_accepted_fixtures_skip_without_a_pack(make_client, fake):
    by_id, _ = results(make_client(checks=only(*AUDIT_CHECKS)))
    assert by_id["P3-UPLOAD-CLEAN"]["message"] == "Not executed: no fixture pack with a clean entry was frozen at setup."
    assert all(by_id[t]["outcome"] == "skip" for t in ACCEPTED_CASES) and uploads(fake) == []


##  ---- Versions, access, determinism, session header, controlled failures ----

def full_client(make_client, tmp_path):
    return make_client(settings=make_settings(tmp_path, FAST_POLL, expectations_path=SHARED_PACK))


def override(fake, match, transform):
    ##  Requests where match(request) is true get transform(status, body) of the fake project's reply.
    def reply(request):
        if not match(request):
            return None
        status, body = fake.project.handle(request.method, request.url.path.removeprefix("/system"),
                                           dict(request.headers), request.content)
        return transform(status, json.loads(json.dumps(body)))
    fake.project_reply = reply


def is_read(kind, user=None):
    def match(request):
        path, token = request.url.path, request.headers.get("authorization", "")
        shaped = ("/audits/" in path) if kind == "audit" else ("/versions/" in path and "/audits/" not in path)
        return request.method == "GET" and shaped and (user is None or token == f"Bearer {user}")
    return match


def test_every_case_but_the_failure_envelope_passes_against_a_conforming_target(make_client, fake, tmp_path):
    by_id, body = results(full_client(make_client, tmp_path))
    assert [t for t, r in by_id.items() if r["outcome"] != "pass"] == ["P3-AUDIT-FAILURE-ENVELOPE"]
    assert by_id["P3-AUDIT-FAILURE-ENVELOPE"]["message"] == \
        "Not executed: no controlled audit failure (PROCESSING_FAILED) was observed in this run."
    assert by_id["P3-DETERMINISM-5"]["message"] == ("Seeded audit 5 matches the first audit: REVIEW_REQUIRED, the "
                                                    "same 8 findings with identical IDs and text, and the same limitations.")
    assert body["coverage"]["component"]["covered"] == 6
    assert body["coverage"]["system"]["covered"] == 7             # every AC except AC-14 (no controlled failure)
    assert sum(1 for _, fields in uploads(fake) if fields.get("file", ("",))[0] == "seeded_issues.xlsx") == 5


@pytest.mark.parametrize("match, transform, test_id, text", [
    (is_read("version", "tok-b"), lambda s, b: (200, {}), "P3-VERSION-FOREIGN", "Expected 403"),
    (is_read("audit", "tok-b"), lambda s, b: (404, b), "P3-AUDIT-FOREIGN", "Expected 403"),
    (is_read("version", "tok-r"), lambda s, b: (403, {"code": "FORBIDDEN", "message": "No.", "requestId": "r1",
                                                       "details": {}}), "P3-VERSION-REVIEWER", "Expected 200"),
    (is_read("audit", "tok-r"), lambda s, b: (s, {**b, "findings": []}), "P3-AUDIT-REVIEWER",
     "findings differs from the owner's read"),
    (lambda r: is_read("version")(r) and "authorization" not in r.headers, lambda s, b: (403, b),
     "P3-VERSION-UNAUTHENTICATED", "Expected 401"),
    (lambda r: is_read("audit")(r) and "authorization" not in r.headers, lambda s, b: (404, b),
     "P3-AUDIT-UNAUTHENTICATED", "Expected 401"),
    (lambda r: r.headers.get("x-test-session", "").startswith("unknown-"),
     lambda s, b: (200, {}), "P3-SESSION-UNKNOWN", "Expected 400"),
])
def test_access_cases_fail_on_the_wrong_answer(make_client, fake, tmp_path, match, transform, test_id, text):
    override(fake, match, transform)
    result = results(full_client(make_client, tmp_path))[0][test_id]
    assert result["outcome"] == "fail" and text in result["message"], result


def test_unknown_and_mismatched_ids_must_not_resolve(make_client, fake, tmp_path):
    ##  A target that ignores unknown IDs or the workbook/version pairing answers 200 instead of 404.
    override(fake, lambda r: r.method == "GET" and "/workbooks/" in r.url.path,
             lambda status, body: (200, {}) if status == 404 else (status, body))
    by_id = results(full_client(make_client, tmp_path))[0]
    for test_id in ("P3-NOT-FOUND", "P3-VERSION-MISMATCH", "P3-AUDIT-MISMATCH", "P3-AUDIT-UNKNOWN"):
        assert by_id[test_id]["outcome"] == "fail" and "Expected 404" in by_id[test_id]["message"], by_id[test_id]
    assert by_id["P3-AUDIT-SEEDED"]["outcome"] == "pass"


@pytest.mark.parametrize("change, test_id, text", [
    (lambda b: b.update(sha256="0" * 64), "P3-SOURCE-HASH", "is not the tester's SHA256"),
    (lambda b: b.update(status="AUDIT_COMPLETE"), "P3-SOURCE-HASH", "but the finished audit is REVIEW_REQUIRED"),
    (lambda b: b.update(fileName="upload.xlsx"), "P3-VERSION-VALID", "fileName 'upload.xlsx', expected 'clean.xlsx'"),
    (lambda b: b.update(latestAuditId=None), "P3-VERSION-VALID", "latestAuditId None"),
    (lambda b: b.update(sourceVersionId="ver-0"), "P3-VERSION-VALID", "sourceVersionId ver-0 is set"),
    (lambda b: b.update(status="REMEDIATED"), "P3-VERSION-VALID", "expected PROCESSING or AUDIT_COMPLETE"),
    (lambda b: b.update(createdAt="yesterday"), "P3-VERSION-VALID", "does not match VersionResponse (createdAt"),
    (lambda b: b.pop("sourceVersionId"), "P3-VERSION-REVIEWER", "does not match VersionResponse (sourceVersionId"),
])
def test_version_reads_report_each_problem(make_client, fake, tmp_path, change, test_id, text):
    def transform(status, body):
        if status == 200:
            change(body)
        return status, body
    override(fake, is_read("version"), transform)
    result = results(full_client(make_client, tmp_path))[0][test_id]
    assert result["outcome"] == "fail" and text in result["message"], result


def test_determinism_catches_changed_prose(make_client, fake, tmp_path):
    seen = []

    def transform(status, body):
        if status == 200 and body["status"] == "REVIEW_REQUIRED" and len(body["findings"]) == 8:
            if body["auditId"] not in seen:
                seen.append(body["auditId"])
            if seen.index(body["auditId"]) >= 2:           # every seeded audit after the first and the reviewer's
                body["findings"][0]["message"] = "Reworded."
        return status, body
    override(fake, is_read("audit", "tok-a"), transform)
    by_id = results(full_client(make_client, tmp_path))[0]
    assert by_id["P3-AUDIT-SEEDED"]["outcome"] == "pass"
    assert by_id["P3-DETERMINISM-3"]["outcome"] == "fail"
    assert by_id["P3-DETERMINISM-3"]["message"] == "Seeded audit 3: findings differs from the first audit."


@pytest.mark.parametrize("failure, outcome, text", [
    ({"completedAt": "2099-01-01T00:00:00Z"}, "pass", "1 controlled audit failure carried"),
    ({"completedAt": "2000-01-01T00:00:00Z"}, "fail", "P3-AUDIT-UNICODE: completedAt is earlier than startedAt"),
    ({"completedAt": None}, "fail", "P3-AUDIT-UNICODE: completedAt is null"),
    ({"error": None}, "fail", "P3-AUDIT-UNICODE: error is null"),
    ({"findings": "keep"}, "fail", "P3-AUDIT-UNICODE: exposes 6 partial findings"),
])
def test_controlled_failure_envelope(make_client, fake, tmp_path, failure, outcome, text):
    def transform(status, body):
        if status == 200 and body["status"] != "PROCESSING" and any(f["worksheet"] == "Café costs"
                                                                    for f in body["findings"]):
            failed = {"status": "PROCESSING_FAILED", "findings": [], "limitations": [],
                      "error": {"code": "AUDIT_FAILED", "message": "No.", "requestId": "r1", "details": {}},
                      **failure}
            if failed["findings"] == "keep":
                failed["findings"] = body["findings"]
            body.update(failed)
        return status, body
    override(fake, is_read("audit"), transform)
    by_id = results(full_client(make_client, tmp_path))[0]
    assert by_id["P3-AUDIT-UNICODE"]["outcome"] == "fail"
    result = by_id["P3-AUDIT-FAILURE-ENVELOPE"]
    assert result["outcome"] == outcome and text in result["message"], result


##  ---- Generated workbooks ----

def test_generated_workbooks_are_deterministic_and_exact():
    assert workbook() == workbook()
    big = workbook(size=25_000_000)
    assert len(big) == 25_000_000
    with zipfile.ZipFile(io.BytesIO(big)) as package:
        assert package.testzip() is None and PADDING_PART in package.namelist()
    assert PADDING_PART not in zipfile.ZipFile(io.BytesIO(workbook())).namelist()
    with pytest.raises(ValueError):
        workbook(size=100)


@pytest.mark.parametrize("worksheets, cells", [(1, 0), (1, 250_001), (51, 3)])
def test_generated_counts_open_in_openpyxl(worksheets, cells):
    book = openpyxl.load_workbook(io.BytesIO(workbook(worksheets=worksheets, cells=cells)))
    assert len(book.worksheets) == worksheets
    assert sum(1 for row in book.worksheets[0].iter_rows() for cell in row if cell.value is not None) == cells


def test_chartsheet_does_not_add_a_worksheet():
    book = openpyxl.load_workbook(io.BytesIO(workbook(worksheets=50, cells=3, chartsheet=True)))
    assert len(book.worksheets) == 50 and [sheet.title for sheet in book.chartsheets] == ["Chart 1"]


def test_structure_protection_is_a_real_password_hash():
    book = openpyxl.load_workbook(io.BytesIO(workbook(protect_structure=True)))
    security = book.security
    assert security.lockStructure and security.workbookAlgorithmName == "SHA-512"
    salt = base64.b64decode(security.workbookSaltValue)
    expected = structure_password_hash("tester-protected", salt, int(security.workbookSpinCount))
    assert base64.b64decode(security.workbookHashValue) == expected


def test_column_letters():
    assert [column(n) for n in (1, 26, 27, 100, 702, 703)] == ["A", "Z", "AA", "CV", "ZZ", "AAA"]
