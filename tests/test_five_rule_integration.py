"""Real-runner five-rule integration, partial WCAG 1.1.1/1.3.1/2.4.2/2.4.4.

Related 2.4.6; REQ-03/04/05/06/15, AC-03/04/05/06/18.
Historical five-rule fixture expectations under the current six-rule runner; no success mocks or six-rule claim.
"""
import hashlib
import json
import sys
import unittest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from accessibility_checker import safe_xml
from accessibility_checker.checks.image_alt_text import check_image_alt_text
from accessibility_checker.checks.worksheet_names import check_worksheet_names
from accessibility_checker.checks.table_headers import check_table_headers
from accessibility_checker.checks.merged_cells import check_merged_cells
from accessibility_checker.checks.link_text import check_link_text
from accessibility_checker.runner import check_workbook
from accessibility_checker.checks.color_contrast import check_color_contrast

FIXTURES = Path(__file__).parent / 'fixtures'
CHECKERS = (check_image_alt_text, check_worksheet_names, check_table_headers,
            check_merged_cells, check_link_text, check_color_contrast)
SMALL = 'TABLE_HEADER: Single cells, one-row regions and one-column lists are not inferred tables under this profile.'
# Authored from the construction recipe and saved OOXML, not checker output.
SEEDED = [
    ('ALT_TEXT', 'NON_TEXT_CONTENT', 'Sheet2', 'xl/drawings/drawing1.xml#id=1', 'REQUIRES_USER_INPUT'),
    ('LINK_TEXT', 'LINK_TEXT', 'Sheet2', 'xl/worksheets/sheet2.xml#cell=E5', 'REQUIRES_USER_INPUT'),
    ('MERGED_CELLS', 'MERGED_CELLS', 'Sheet2', 'xl/worksheets/sheet2.xml#merge=A5:C5', 'MANUAL_REVIEW'),
    ('TABLE_HEADER', 'TABLE_STRUCTURE', 'Data', 'xl/tables/table1.xml#tableId=1', 'REQUIRES_USER_INPUT'),
    ('WORKSHEET_NAME', 'WORKSHEET_NAME', 'Sheet2', 'xl/worksheets/sheet2.xml#sheetId=2', 'REQUIRES_USER_INPUT'),
]


class FiveRuleIntegrationTests(unittest.TestCase):
    def verify(self, seeded):
        path = FIXTURES / ('five_rule_seeded.xlsx' if seeded else 'five_rule_clean.xlsx')
        before = path.read_bytes()
        source = hashlib.sha256(before).hexdigest()
        book = load_workbook(path)
        try:
            self.assertEqual(book.worksheets[1].sheet_state, 'hidden')
            self.assertEqual(book['Archive'].sheet_state, 'veryHidden')
            self.assertEqual(len(book['Data'].tables), 1)
        finally:
            book.close()
        with ZipFile(path) as archive:
            props = safe_xml.fromstring(archive.read('xl/drawings/drawing1.xml')).find('.//{*}cNvPr')
            self.assertEqual(props.get('id'), '1')
            self.assertEqual(props.get('title'), None if seeded else 'Scores')
            table = safe_xml.fromstring(archive.read('xl/tables/table1.xml'))
            self.assertEqual((table.get('id'), table.get('ref'), table.get('headerRowCount')),
                             ('1', 'A1:B2', '0' if seeded else '1'))
        original_profile = sys.getprofile()
        calls = []
        def trace(frame, event, arg):
            if event == 'call' and '/accessibility_checker/' in frame.f_code.co_filename:
                calls.append((frame.f_code.co_filename, frame.f_code))
        try:
            sys.setprofile(trace)
            result = check_workbook(path)
        finally:
            sys.setprofile(original_profile)
        self.assertEqual([code for _, code in calls if code in {c.__code__ for c in CHECKERS}],
                         [c.__code__ for c in CHECKERS])
        self.assertFalse(any(any(name in filename for name in (
            'workbook_title.py', 'legacy_table_headers.py'))
            for filename, _ in calls))
        fields = ('ruleId', 'category', 'worksheet', 'target', 'classification')
        self.assertEqual([tuple(item[field] for field in fields) for item in result], SEEDED if seeded else [])
        self.assertEqual(result, sorted(result, key=lambda f: (f['ruleId'], f['worksheet'], f['target'])))
        self.assertEqual(result.limitations, ['COLOR_CONTRAST: Missing fill is screened as white.', SMALL])
        self.assertEqual(result.limitations, sorted(set(result.limitations)))
        for item in result:
            self.assertEqual(set(item), set(fields) | {'findingId', 'message', 'recommendation'})
            self.assertTrue(all(type(value) is str for value in item.values()))
            for key in ('message', 'recommendation'):
                self.assertTrue(item[key].strip())
                self.assertLessEqual(len(item[key]), 500)
            payload = [source, '2.0.0', item['ruleId'], item['worksheet'], item['target']]
            self.assertEqual(item['findingId'], hashlib.sha256(json.dumps(
                payload, ensure_ascii=False, separators=(',', ':')).encode('utf-8')).hexdigest())
        for _ in range(4):  # Five executions total, including traced execution.
            repeated = check_workbook(path)
            self.assertEqual(repeated, result)
            self.assertEqual(repeated.limitations, result.limitations)
        combined, limitations = [], []
        for checker in CHECKERS:
            direct = checker(path)
            combined.extend(direct)
            limitations.extend(getattr(direct, 'limitations', []))
        self.assertEqual(sorted(combined, key=lambda f: (f['ruleId'], f['worksheet'], f['target'])), result)
        self.assertEqual(sorted(set(limitations)), result.limitations)
        self.assertEqual(path.read_bytes(), before)

    def test_clean_within_five_rule_scope(self):
        self.verify(False)

    def test_seeded_all_five_rules_including_hidden_and_both_classes(self):
        self.verify(True)
