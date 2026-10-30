"""Partial 1.4.1/1.4.3 screening; REQ-03/04/05/06, AC-03/04/05/06."""

import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from accessibility_checker.checks.color_contrast import check_color_contrast


class ColorContrastTests(unittest.TestCase):
    def build(self, colors, conditional=False):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / "color.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Data"
        for row, (color, size, bold, value) in enumerate(colors, 1):
            cell = sheet.cell(row, 1, value)
            cell.font = Font(color=color, sz=size, b=bold)
        if conditional:
            sheet.conditional_formatting.add(
                "A1:A2", FormulaRule(formula=["TRUE"], fill=PatternFill("solid", fgColor="FF0000"))
            )
            sheet.conditional_formatting.add(
                "A1:A2",
                ColorScaleRule(
                    start_type="min", start_color="FFFFFF", end_type="max", end_color="FF0000"
                ),
            )
            sheet.conditional_formatting.add("A3", FormulaRule(formula=["TRUE"]))
        book.save(path)
        book.close()
        return path

    def test_thresholds_gray_large_and_combined_meaning(self):
        path = self.build(
            [
                ("777777", 11, False, "a"),
                ("777777", 18, False, "b"),
                ("959595", 18, False, "c"),
                ("949494", 14, True, "d"),
                ("FF0000", 11, False, "e"),
                ("000000", 11, False, 0),
            ]
        )
        before = path.read_bytes()
        result = check_color_contrast(path)
        self.assertEqual(
            [f["target"].split("#")[1] for f in result], ["cell=A1", "cell=A3", "cell=A5"]
        )
        self.assertTrue(
            all(
                f["classification"] == "MANUAL_REVIEW" and f["category"] == "COLOR_OR_CONTRAST"
                for f in result
            )
        )
        self.assertIn("Nonneutral", result[-1]["message"])
        self.assertIn("below", result[-1]["message"])
        self.assertEqual(result, check_color_contrast(path))
        self.assertEqual(before, path.read_bytes())

    def test_conditional_dedup_suppression_and_nonvisual(self):
        path = self.build([("FF0000", 11, False, "a")] * 3, True)
        result = check_color_contrast(path)
        self.assertEqual(
            [f["target"].split("#")[1] for f in result], ["cell=A3", "conditional=A1:A2"]
        )
        self.assertTrue(any("not evaluated" in x for x in result.limitations))

    def test_formula_no_cache_screen_and_assumption(self):
        path = self.build([("FFFFFF", 11, False, "=1"), ("FF0000", 11, False, "=1")])
        result = check_color_contrast(path)
        self.assertEqual([f["target"].split("#")[1] for f in result], ["cell=A2"])
        self.assertNotIn("below", result[0]["message"])
        self.assertTrue(any("not calculated" in x for x in result.limitations))
