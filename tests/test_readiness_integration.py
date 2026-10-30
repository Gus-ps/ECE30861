"""Readiness integration for partial 1.1.1 Non-text Content, 2.4.2 Page Titled
and related 2.4.6 Headings and Labels; REQ-03/04/05/06/15, AC-03/04/05/06/18.
Generated XLSX and synthetic OOXML checks, not Excel UI or remediation evidence.
"""

import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

from accessibility_checker.runner import check_workbook


class ReadinessIntegrationTests(unittest.TestCase):
    def setUp(self):
        # Preserve this historical two-rule acceptance profile.
        for name in (
            "check_table_headers",
            "check_merged_cells",
            "check_link_text",
            "check_color_contrast",
        ):
            mocked = patch("accessibility_checker.runner." + name, return_value=[])
            mocked.start()
            self.addCleanup(mocked.stop)

    def assert_common_contract(self, path, result):
        before = path.read_bytes()
        source = hashlib.sha256(before).hexdigest()
        self.assertEqual({f["ruleId"] for f in result}, {"ALT_TEXT", "WORKSHEET_NAME"})
        self.assertEqual(
            result, sorted(result, key=lambda f: (f["ruleId"], f["worksheet"], f["target"]))
        )
        self.assertEqual(result.limitations, sorted(set(result.limitations)))
        self.assertEqual(len({f["findingId"] for f in result}), len(result))
        for finding in result:
            self.assertEqual(
                set(finding),
                {
                    "findingId",
                    "ruleId",
                    "category",
                    "worksheet",
                    "target",
                    "classification",
                    "message",
                    "recommendation",
                },
            )
            identity = json.dumps(
                [source, "2.0.0", finding["ruleId"], finding["worksheet"], finding["target"]],
                ensure_ascii=False,
                separators=(",", ":"),
            ).encode("utf-8")
            self.assertEqual(finding["findingId"], hashlib.sha256(identity).hexdigest())
        repeated = check_workbook(path)
        self.assertEqual(result, repeated)
        self.assertEqual(result.limitations, repeated.limitations)
        self.assertEqual(path.read_bytes(), before)

    def test_generated_complete_xlsx_supports_required_safe_reference_families(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "safe-references.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "Sheet1"
            data = workbook.create_sheet("O'Brien")
            data.append([10, "Item"])
            data.append([20, "Other"])
            first["A1"] = 1
            first["A2"] = 2
            first["B1"] = "=SUM('O''Brien'!$A$1:$A$2)"
            first["B2"] = '=IF(A1>0,"Sheet1!A1", "literal")'
            workbook.defined_names.add(DefinedName("InputRange", attr_text="'O''Brien'!$A$1:$A$2"))
            data.defined_names.add(DefinedName("LocalInput", attr_text="'Sheet1'!$A$1"))
            data["B3"] = "=SUM(LocalInput,Sheet1!A1)"
            validation = DataValidation(type="list", formula1="InputRange")
            first.add_data_validation(validation)
            validation.add("C1")
            first.conditional_formatting.add(
                "A1:A2",
                FormulaRule(
                    formula=["A1>0"],
                    fill=PatternFill(fill_type="solid", fgColor="FFFF0000"),
                ),
            )
            first["A4"] = "Go to data"
            first["A4"].hyperlink = Hyperlink(ref="A4", location="'O''Brien'!A1")
            chart = BarChart()
            chart.add_data(Reference(data, min_col=1, min_row=1, max_row=2))
            chart_sheet = workbook.create_chartsheet("Sheet3")
            chart_sheet.add_chart(chart)
            chart_sheet.sheet_state = "hidden"
            workbook.save(path)  # Fixture creation only; the checker never saves.
            workbook.close()
            before = path.read_bytes()
            loaded = load_workbook(path)
            try:
                self.assertEqual(loaded["Sheet1"]["B1"].value, "=SUM('O''Brien'!$A$1:$A$2)")
                self.assertIn("InputRange", loaded.defined_names)
                self.assertEqual(loaded.sheetnames, ["Sheet1", "O'Brien", "Sheet3"])
            finally:
                loaded.close()
            result = check_workbook(path)
            self.assertTrue(
                all(f["classification"] == "REQUIRES_USER_INPUT" for f in result),
                result.limitations,
            )
            self.assertEqual(
                {f["worksheet"] for f in result if f["ruleId"] == "WORKSHEET_NAME"},
                {"Sheet1", "Sheet3"},
            )
            self.assertEqual(result.limitations, [])
            self.assert_common_contract(path, result)
            self.assertEqual(before, path.read_bytes())

    def test_complete_fixture_shared_drawing_integrates_both_classifications(self):
        fixture = Path(__file__).parent / "fixtures" / "migration_complete.xlsx"
        original = fixture.read_bytes()
        main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        doc = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        pkg = "http://schemas.openxmlformats.org/package/2006/relationships"
        with ZipFile(fixture) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        sheet = ET.fromstring(parts["xl/worksheets/sheet2.xml"])
        ET.SubElement(sheet, f"{{{main}}}drawing", {f"{{{doc}}}id": "shared"})
        parts["xl/worksheets/sheet2.xml"] = ET.tostring(sheet)
        parts["xl/worksheets/_rels/sheet2.xml.rels"] = (
            f'<Relationships xmlns="{pkg}"><Relationship Id="shared" Type="{doc}/drawing" '
            'Target="../drawings/drawing1.xml"/></Relationships>'
        ).encode("utf-8")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "shared.xlsx"
            with ZipFile(path, "w") as archive:
                for name, data in parts.items():
                    archive.writestr(name, data)
            result = check_workbook(path)
            shared = [
                f
                for f in result
                if f["ruleId"] == "ALT_TEXT" and f["target"] == "xl/drawings/drawing1.xml#id=1"
            ]
            self.assertEqual({f["worksheet"] for f in shared}, {"Sheet01", "sheet2"})
            self.assertTrue(all(f["classification"] == "MANUAL_REVIEW" for f in shared))
            names = [f for f in result if f["ruleId"] == "WORKSHEET_NAME"]
            self.assertEqual(len(names), 3)
            self.assertTrue(all(f["classification"] == "MANUAL_REVIEW" for f in names))
            self.assertTrue(any("shared drawing ownership" in s for s in result.limitations))
            self.assert_common_contract(path, result)
        self.assertEqual(fixture.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
