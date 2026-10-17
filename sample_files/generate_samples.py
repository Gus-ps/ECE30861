"""
Generates sample files for testing validity_check.py.

Each file violates exactly one requirement (except valid.xlsx, which meets all
of them), so the output of the validity checker can be compared against the
expected result listed below.

    valid.xlsx              -> valid
    too_large.xlsx          -> over the 25 MB file size limit
    too_many_sheets.xlsx    -> 51 worksheets (limit is 50)
    too_many_cells.xlsx     -> 250,001 nonempty cells (limit is 250,000)
    corrupted.xlsx          -> damaged archive (simulates a broken download)
    encrypted.xlsx          -> password required to open (password: "password")
    protected_sheet.xlsx    -> worksheet protected with a password
    protected_workbook.xlsx -> workbook structure protected with a password
    wrong_file_type.csv     -> not a .xlsx file

Requirements (not needed by the validity checker itself):
    pip install openpyxl msoffcrypto-tool

Usage:
    python generate_samples.py
"""

import io
import os
import random
import string

import msoffcrypto
from openpyxl import Workbook

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
PASSWORD = "password"


def path(name):
    return os.path.join(OUTPUT_DIR, name)


def make_basic_workbook():
    """A small, normal workbook: a header row and a few rows of data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Department", "Quarter", "Spending"])
    rows = [("Admissions", "Q1", 12000), ("Library", "Q1", 8500),
            ("Athletics", "Q2", 15300), ("Housing", "Q2", 21000)]
    for row in rows:
        ws.append(row)
    ws["D1"] = "Total"
    ws["D2"] = "=SUM(C2:C5)"
    return wb


def make_valid():
    make_basic_workbook().save(path("valid.xlsx"))


def make_too_large():
    # 200,000 cells (under the cell limit) of random text that compresses
    # poorly, so the saved file ends up over 25 MB.
    rng = random.Random(0)
    alphabet = string.ascii_letters + string.digits
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    for _ in range(50_000):
        ws.append(["".join(rng.choices(alphabet, k=200)) for _ in range(4)])
    wb.save(path("too_large.xlsx"))


def make_too_many_sheets():
    wb = make_basic_workbook()
    for i in range(2, 52):
        wb.create_sheet(f"Sheet{i}")["A1"] = f"Sheet {i}"
    wb.save(path("too_many_sheets.xlsx"))


def make_too_many_cells():
    # 50,001 rows x 5 columns = 250,005 cells; trim the last row to land on
    # exactly 250,001.
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    for r in range(50_000):
        ws.append([r, r + 1, r + 2, r + 3, r + 4])
    ws.append([1])
    wb.save(path("too_many_cells.xlsx"))


def make_corrupted():
    # Save a valid workbook, then cut off the end of the file, the way an
    # interrupted download or transfer would.
    buffer = io.BytesIO()
    make_basic_workbook().save(buffer)
    data = buffer.getvalue()
    with open(path("corrupted.xlsx"), "wb") as f:
        f.write(data[: len(data) // 2])


def make_encrypted():
    buffer = io.BytesIO()
    make_basic_workbook().save(buffer)
    buffer.seek(0)
    office_file = msoffcrypto.OfficeFile(buffer)
    with open(path("encrypted.xlsx"), "wb") as f:
        office_file.encrypt(PASSWORD, f)


def make_protected_sheet():
    wb = make_basic_workbook()
    wb.active.protection.set_password(PASSWORD)
    wb.active.protection.enable()
    wb.save(path("protected_sheet.xlsx"))


def make_protected_workbook():
    wb = make_basic_workbook()
    wb.security.workbookPassword = PASSWORD
    wb.security.lockStructure = True
    wb.save(path("protected_workbook.xlsx"))


def make_wrong_file_type():
    with open(path("wrong_file_type.csv"), "w", newline="") as f:
        f.write("Department,Quarter,Spending\nAdmissions,Q1,12000\n")


def main():
    for make in (make_valid, make_too_large, make_too_many_sheets,
                 make_too_many_cells, make_corrupted, make_encrypted,
                 make_protected_sheet, make_protected_workbook,
                 make_wrong_file_type):
        make()
        print(f"Created {make.__name__[len('make_'):]}")


if __name__ == "__main__":
    main()
